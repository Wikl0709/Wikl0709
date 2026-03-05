import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
from typing import List, Dict, Optional


class ExcelStyleFormatter:
    """
    Excel 样式格式化器
    用于对 Excel 文件进行排序、着色和样式设置
    """
    
    # 默认排序列
    DEFAULT_SORT_COLUMNS = [
        'Domain Classifier',
        'Intent Mapping',
        'Intent Recognition',
        'Slot Extraction',
        'Tool Execution',
        'Memory Retrieval',
        'Knowledge Retrieval',
        'Final Answer'
    ]
    
    # 默认配色方案
    DEFAULT_COLOR_SCHEME = {
        'domain_pass': '00B050',      # 深绿
        'domain_pass*': '92D050',     # 中绿
        'domain_pass**': 'C5E1A5',    # 浅绿
        'slot_pass': '00B050',        # 深绿
        'slot_pass*': 'FFB74D',       # 橙色
        'slot_fail': 'FF0000',        # 红色
        'other_pass': '00B050',       # 深绿
        'other_fail': 'FF0000',       # 红色
    }
    
    # 预设配色方案
    COLOR_SCHEMES = {
        'default': DEFAULT_COLOR_SCHEME,
        'morandi': {
            'domain_pass': '4A7C59',
            'domain_pass*': '7FA99B',
            'domain_pass**': 'B4C7C9',
            'slot_pass': '4A7C59',
            'slot_pass*': 'E8B4B8',
            'slot_fail': 'C97064',
            'other_pass': '4A7C59',
            'other_fail': 'C97064',
        },
        'fresh': {
            'domain_pass': '2D5016',
            'domain_pass*': '6B8E23',
            'domain_pass**': '9ACD32',
            'slot_pass': '2D5016',
            'slot_pass*': 'FFB347',
            'slot_fail': 'CD5C5C',
            'other_pass': '2D5016',
            'other_fail': 'CD5C5C',
        },
        'gentle': {
            'domain_pass': '3D7068',
            'domain_pass*': '5F9EA0',
            'domain_pass**': '98D8C8',
            'slot_pass': '3D7068',
            'slot_pass*': 'F4A460',
            'slot_fail': 'E07B7B',
            'other_pass': '3D7068',
            'other_fail': 'E07B7B',
        },
        'elegant': {
            'domain_pass': '4F7942',
            'domain_pass*': '87A96B',
            'domain_pass**': 'C5E1A5',
            'slot_pass': '4F7942',
            'slot_pass*': 'FFB74D',
            'slot_fail': 'E57373',
            'other_pass': '4F7942',
            'other_fail': 'E57373',
        }
    }
    
    def __init__(self, 
                 sort_columns: Optional[List[str]] = None,
                 color_scheme: str = 'default',
                 custom_colors: Optional[Dict[str, str]] = None,
                 border_style: str = 'thick',
                 border_color: str = 'FFFFFF',
                 font_name: str = 'Calibri',
                 font_size: int = 11):
        """
        初始化格式化器
        
        Args:
            sort_columns: 需要排序的列名列表
            color_scheme: 配色方案名称 ('default', 'morandi', 'fresh', 'gentle', 'elegant')
            custom_colors: 自定义颜色字典
            border_style: 边框样式 ('thin', 'medium', 'thick', 'double')
            border_color: 边框颜色 (十六进制)
            font_name: 字体名称
            font_size: 字体大小
        """
        self.sort_columns = sort_columns or self.DEFAULT_SORT_COLUMNS
        
        # 设置配色方案
        if custom_colors:
            self.color_scheme = custom_colors
        elif color_scheme in self.COLOR_SCHEMES:
            self.color_scheme = self.COLOR_SCHEMES[color_scheme]
        else:
            print(f"⚠️ 未知的配色方案 '{color_scheme}'，使用默认方案")
            self.color_scheme = self.DEFAULT_COLOR_SCHEME
        
        # 创建 PatternFill 对象
        self.colors = {
            key: PatternFill(start_color=color, end_color=color, fill_type='solid')
            for key, color in self.color_scheme.items()
        }
        
        # 边框样式
        self.border = Border(
            left=Side(style=border_style, color=border_color),
            right=Side(style=border_style, color=border_color),
            top=Side(style=border_style, color=border_color),
            bottom=Side(style=border_style, color=border_color)
        )
        
        # 字体样式
        self.data_font = Font(name=font_name, size=font_size)
        self.header_font = Font(name=font_name, size=font_size, bold=True)
        
        # 对齐方式
        self.center_alignment = Alignment(horizontal='center', vertical='center')
    
    @staticmethod
    def create_sort_key(value):
        """
        创建排序键值
        
        Args:
            value: 单元格值
            
        Returns:
            排序键值 (0-4)
        """
        if pd.isna(value):
            return 3
        
        value_str = str(value).lower()
        
        if 'pass**' in value_str:
            return 2
        elif 'pass*' in value_str:
            return 1
        elif 'pass' in value_str:
            return 0
        elif 'fail' in value_str:
            return 4
        else:
            return 3
    
    def sort_dataframe(self, df: pd.DataFrame, 
                      primary_sort_columns: Optional[List[str]] = None) -> pd.DataFrame:
        """
        对 DataFrame 进行排序
        
        Args:
            df: 输入的 DataFrame
            primary_sort_columns: 主排序列（如 ['domain', 'secondary']）
            
        Returns:
            排序后的 DataFrame
        """
        # 为每个需要排序的列创建排序键
        for col in self.sort_columns:
            if col in df.columns:
                df[f'{col}_sort_key'] = df[col].apply(self.create_sort_key)
        
        # 创建排序键列的列表
        sort_key_columns = [f'{col}_sort_key' for col in self.sort_columns if col in df.columns]
        
        # 确定最终的排序列
        if primary_sort_columns:
            final_sort_columns = [col for col in primary_sort_columns if col in df.columns] + sort_key_columns
        else:
            final_sort_columns = sort_key_columns
        
        # 执行排序
        df_sorted = df.sort_values(by=final_sort_columns, ascending=True)
        
        # 删除临时的排序键列
        df_sorted = df_sorted.drop(columns=sort_key_columns)
        
        # 重置索引
        df_sorted = df_sorted.reset_index(drop=True)
        
        return df_sorted
    
    def apply_cell_color(self, cell, column_name: str):
        """
        为单元格应用颜色
        
        Args:
            cell: openpyxl 单元格对象
            column_name: 列名
        """
        value = str(cell.value).lower() if cell.value else ''
        
        # Domain Classifier 列
        if column_name == 'Domain Classifier':
            if 'pass**' in value:
                cell.fill = self.colors['domain_pass**']
            elif 'pass*' in value:
                cell.fill = self.colors['domain_pass*']
            elif 'pass' in value:
                cell.fill = self.colors['domain_pass']
        
        # Slot Extraction 列
        elif column_name == 'Slot Extraction':
            if 'pass*' in value:
                cell.fill = self.colors['slot_pass*']
            elif 'pass' in value:
                cell.fill = self.colors['slot_pass']
            elif 'fail' in value:
                cell.fill = self.colors['slot_fail']
        
        # 其他列
        elif column_name in ['Intent Mapping', 'Intent Recognition', 
                            'Tool Execution', 'Memory Retrieval', 'Knowledge Retrieval', 'Final Answer']:
            if 'pass' in value:
                cell.fill = self.colors['other_pass']
            elif 'fail' in value:
                cell.fill = self.colors['other_fail']
    
    def apply_styles(self, workbook_path: str, df: pd.DataFrame, 
                    auto_adjust_width: bool = True, max_width: int = 50):
        """
        为 Excel 文件应用样式
        
        Args:
            workbook_path: Excel 文件路径
            df: DataFrame（用于获取行列数）
            auto_adjust_width: 是否自动调整列宽
            max_width: 最大列宽
        """
        # 加载工作簿
        wb = load_workbook(workbook_path)
        ws = wb.active
        
        # 获取列索引
        column_indices = {col: idx for idx, col in enumerate(df.columns, start=1)}
        
        # 为所有单元格设置边框和字体
        for row in ws.iter_rows(min_row=1, max_row=len(df) + 1, 
                               min_col=1, max_col=len(df.columns)):
            for cell in row:
                cell.border = self.border
                
                if cell.row == 1:  # 表头
                    cell.font = self.header_font
                    cell.alignment = self.center_alignment
                else:  # 数据行
                    cell.font = self.data_font
        
        # 为单元格着色
        for row_idx in range(2, len(df) + 2):  # 从第2行开始（第1行是表头）
            for col_name, col_idx in column_indices.items():
                if col_name in self.sort_columns:
                    cell = ws.cell(row=row_idx, column=col_idx)
                    self.apply_cell_color(cell, col_name)
                    
                    # 重新应用边框和字体
                    cell.border = self.border
                    cell.font = self.data_font
        
        # 自动调整列宽
        if auto_adjust_width:
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, max_width)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # 保存文件
        wb.save(workbook_path)
    
    def format_excel(self, 
                    input_path: str, 
                    output_path: str,
                    primary_sort_columns: Optional[List[str]] = None,
                    auto_adjust_width: bool = True,
                    max_width: int = 50) -> pd.DataFrame:
        """
        完整的 Excel 格式化流程
        
        Args:
            input_path: 输入 Excel 文件路径
            output_path: 输出 Excel 文件路径
            primary_sort_columns: 主排序列（如 ['domain', 'secondary']）
            auto_adjust_width: 是否自动调整列宽
            max_width: 最大列宽
            
        Returns:
            格式化后的 DataFrame
        """
        print(f"📖 读取文件: {input_path}")
        
        # 读取 Excel 文件
        df = pd.read_excel(input_path)
        
        print(f"📊 数据行数: {len(df)}")
        print(f"📋 数据列数: {len(df.columns)}")
        
        # 排序
        print("🔄 正在排序...")
        df_sorted = self.sort_dataframe(df, primary_sort_columns)
        
        # 保存到新文件
        print(f"💾 保存文件: {output_path}")
        df_sorted.to_excel(output_path, index=False)
        
        # 应用样式
        print("🎨 应用样式...")
        self.apply_styles(output_path, df_sorted, auto_adjust_width, max_width)
        
        print(f"✅ 文件已保存并着色完成: {output_path}")
        print(f"✅ 配色方案: {self._get_scheme_name()}")
        print(f"✅ 边框样式: 白色粗边框")
        print(f"✅ 字体: {self.data_font.name}")
        
        return df_sorted
    
    def _get_scheme_name(self) -> str:
        """获取当前配色方案名称"""
        for name, scheme in self.COLOR_SCHEMES.items():
            if scheme == self.color_scheme:
                return name
        return "自定义"
    
    @classmethod
    def list_color_schemes(cls):
        """列出所有可用的配色方案"""
        print("\n🎨 可用的配色方案:")
        print("=" * 60)
        for name, scheme in cls.COLOR_SCHEMES.items():
            print(f"\n方案名称: {name}")
            print(f"  Domain Pass:   #{scheme['domain_pass']}")
            print(f"  Domain Pass*:  #{scheme['domain_pass*']}")
            print(f"  Domain Pass**: #{scheme['domain_pass**']}")
            print(f"  Slot Pass:     #{scheme['slot_pass']}")
            print(f"  Slot Pass*:    #{scheme['slot_pass*']}")
            print(f"  Slot Fail:     #{scheme['slot_fail']}")
            print(f"  Other Pass:    #{scheme['other_pass']}")
            print(f"  Other Fail:    #{scheme['other_fail']}")


# 使用示例
if __name__ == '__main__':
    # ============ 示例 1: 使用默认配置 ============
    formatter = ExcelStyleFormatter()
    
    input_file = r"D:\00work\AIPC\Intent\20260112_tool\localsolutionLogTool 2\localsolutionLogTool\Results-en-0112\test_results_local_solution_0112_slot_20260112_Auto_Judge_Final.xlsx"
    output_file = r"D:\00work\AIPC\Intent\20260112_tool\localsolutionLogTool 2\localsolutionLogTool\Results-en-0112\test_results_local_solution_0112_slot_20260112_Auto_Judge_Final_coloerd.xlsx"
    
    df_result = formatter.format_excel(
        input_path=input_file,
        output_path=output_file,
        primary_sort_columns=['domain', 'secondary']
    )
    
    # ============ 示例 2: 使用莫兰迪配色方案 ============
    # formatter_morandi = ExcelStyleFormatter(color_scheme='morandi')
    # df_result = formatter_morandi.format_excel(input_file, output_file)
    
    # ============ 示例 3: 自定义配色 ============
    # custom_colors = {
    #     'domain_pass': '4A7C59',
    #     'domain_pass*': '7FA99B',
    #     'domain_pass**': 'B4C7C9',
    #     'slot_pass': '4A7C59',
    #     'slot_pass*': 'E8B4B8',
    #     'slot_fail': 'C97064',
    #     'other_pass': '4A7C59',
    #     'other_fail': 'C97064',
    # }
    # formatter_custom = ExcelStyleFormatter(custom_colors=custom_colors)
    # df_result = formatter_custom.format_excel(input_file, output_file)
    
    # ============ 查看所有配色方案 ============
    # ExcelStyleFormatter.list_color_schemes()
